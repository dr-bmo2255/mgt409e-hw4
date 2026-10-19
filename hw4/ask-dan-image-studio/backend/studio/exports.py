"""Phase 3 export: approved individual PNGs + mapping JSON/CSV + refreshed XLSX/TXT + review/cost report + ZIP.
Only APPROVED images are exported. Remaining missing / needs-review assets are listed separately."""

from __future__ import annotations

import csv
import json
import shutil
import time
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from . import config, state
from .engine import is_mock_derived, manifest
from .inventory import find_inputs
from .qa import secrets_leak_scan


def _rows():
    m = manifest()
    states = state.all_asset_states()
    assets = []
    for a in m["assets"]:
        s = states[a["asset_id"]]
        approved = config.APPROVED_DIR / f"{a['asset_id']}.png"
        mock = bool(s.get("approved_version")) and is_mock_derived(a["asset_id"], s["approved_version"])
        assets.append({**a, "studio_status": s["status"], "approved_version": s.get("approved_version"), "mock_approval": mock,
                       "approved_file": approved.name if (s["status"] == "APPROVED" and approved.exists() and not mock) else None,
                       "cost_usd": round(s.get("cost_usd", 0) or 0, 4), "versions": s.get("version_count", 0)})
    return m, assets


def export_all() -> dict:
    m, assets = _rows()
    stamp = time.strftime("%Y%m%d-%H%M%S")
    out = config.EXPORT_DIR / f"ask_dan_export_{stamp}"
    (out / "approved_png").mkdir(parents=True)
    by_id = {a["asset_id"]: a for a in assets}

    exported, custom = [], []
    for a in assets:
        if a["approved_file"]:
            sub = "custom_png" if a.get("custom") else "approved_png"
            (out / sub).mkdir(exist_ok=True)
            shutil.copy2(config.APPROVED_DIR / a["approved_file"], out / sub / a["canonical_filename"])
            (custom if a.get("custom") else exported).append(a["canonical_filename"])

    # Mapping JSON / CSV (product → Dan asset, with current status and file)
    assets_catalogue = [a for a in assets if not a.get("custom")]
    mapping = []
    for p in m["products"]:
        a = by_id[p["dan_asset_id"]]
        mapping.append({"product_id": p["product_id"], "product_name": p["name"], "product_photo": p["photo_file"],
                        "dan_group": p["dan_group"], "dan_category": p["dan_category"], "dan_asset_id": a["asset_id"],
                        "status": a["studio_status"],
                        "dan_file": a["approved_file"] and a["canonical_filename"] or "; ".join(a["existing_files"]) or None})
    (out / "product_dan_mapping.json").write_text(json.dumps(mapping, indent=2))
    with open(out / "product_dan_mapping.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(mapping[0].keys()))
        w.writeheader()
        w.writerows(mapping)

    # Refreshed workbook: original sheets kept, checklist statuses updated, studio sheet added
    src = find_inputs()["audit_xlsx"]
    wb = openpyxl.load_workbook(src) if src else openpyxl.Workbook()
    if "Dan Asset Checklist" in wb.sheetnames:
        ws = wb["Dan Asset Checklist"]
        hdr = [c.value for c in ws[1]]
        if "Studio status" not in hdr:
            ws.cell(1, len(hdr) + 1, "Studio status").font = Font(bold=True)
            ws.cell(1, len(hdr) + 2, "Studio file").font = Font(bold=True)
            hdr += ["Studio status", "Studio file"]
        fcol = hdr.index("Suggested output filename") + 1 if "Suggested output filename" in hdr else None
        for row in ws.iter_rows(min_row=2):
            fn = row[fcol - 1].value if fcol else None
            aid = (fn or "").replace(".png", "")
            a = by_id.get(aid) or next((x for x in assets if x["canonical_filename"] == fn), None)
            if a:
                ws.cell(row[0].row, hdr.index("Studio status") + 1, a["studio_status"])
                ws.cell(row[0].row, hdr.index("Studio file") + 1, a["approved_file"] and a["canonical_filename"] or "; ".join(a["existing_files"]))
    if "Studio Status" in wb.sheetnames:
        del wb["Studio Status"]
    ws = wb.create_sheet("Studio Status", 0)
    cols = ["asset_id", "group", "category", "product_count", "studio_status", "approved_version", "canonical_filename",
            "existing_files", "shield_slug", "versions", "cost_usd"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="00356B")
    for a in assets:
        ws.append([a[k] if k != "existing_files" else "; ".join(a[k]) for k in cols])
    wb.save(out / "ask_dan_data2_audit_updated.xlsx")

    # Refreshed TXT checklist
    lines = ["ASK DAN — CATALOGUE CHECKLIST (exported by Ask Dan Image Studio)", time.strftime("%Y-%m-%d %H:%M"), ""]
    for title, pred in [("APPROVED NEW ASSETS", lambda a: a["studio_status"] == "APPROVED"),
                        ("EXISTING ASSETS (reused)", lambda a: a["studio_status"] == "EXISTING_REVIEWED"),
                        ("NEEDS REVIEW", lambda a: a["studio_status"] in ("PENDING_REVIEW", "REVISION_REQUESTED")),
                        ("STILL MISSING / NOT APPROVED", lambda a: a["studio_status"] in ("MISSING", "QUEUED", "REJECTED", "FAILED", "GENERATING"))]:
        group = [a for a in assets_catalogue if pred(a)]
        lines += [f"{title} ({len(group)})", "=" * len(title)]
        lines += [f"[{'x' if a['studio_status'] in ('APPROVED', 'EXISTING_REVIEWED') else ' '}] {a['group']}: {a['category']} | "
                  f"{a['canonical_filename'] if a['studio_status'] != 'EXISTING_REVIEWED' else '; '.join(a['existing_files'])} | "
                  f"{a['product_count']} product(s) | {a['studio_status']}" for a in group] + [""]
    customs = [a for a in assets if a.get("custom")]
    if customs:
        lines += [f"CUSTOM ONE-OFF IMAGES (from chat, not catalogue-scoped) ({len(customs)})", "=" * 40]
        lines += [f"[{'x' if a['studio_status'] == 'APPROVED' else ' '}] {a['display_name']} | {a['canonical_filename']} | {a['studio_status']}" for a in customs] + [""]
    (out / "ask_dan_checklist_updated.txt").write_text("\n".join(lines))

    # Review / cost report
    total = state.total_spent()
    with state.db() as c:
        versions = [dict(r) for r in c.execute("SELECT asset_id, version, kind, provider, model, cost_usd, cost_is_estimate, decision, qa_json FROM versions ORDER BY asset_id, version")]
        fb = [dict(r) for r in c.execute("SELECT asset_id, version, kind, text FROM feedback WHERE text<>'' ORDER BY id")]
    rep = ["# Ask Dan Image Studio — review & cost report", f"Exported {time.strftime('%Y-%m-%d %H:%M')}", "",
           f"- Approved new catalogue assets exported: **{len(exported)}** / {sum(1 for a in assets_catalogue if a['status'] == 'MISSING')} scoped",
           f"- Approved custom one-off images exported: **{len(custom)}**",
           f"- Total spend: **${total:.2f}** ({'includes estimates' if any(v['cost_is_estimate'] and v['provider'] == 'openai' for v in versions) else 'actual reported usage'})",
           "", "## Versions", "| Asset | v | kind | provider | cost | decision | QA |", "|---|---|---|---|---|---|---|"]
    for v in versions:
        qa = json.loads(v["qa_json"] or "{}")
        rep.append(f"| {v['asset_id']} | {v['version']} | {v['kind']} | {v['provider']} | ${v['cost_usd']:.3f}{'*' if v['cost_is_estimate'] and v['provider'] == 'openai' else ''} | "
                   f"{v['decision']} | {'OK' if qa.get('ok', True) else '; '.join(qa.get('errors', []))} |")
    rep += ["", "\\* estimated cost (provider did not report usage)", "", "## Reviewer feedback"]
    rep += [f"- {f['asset_id']} v{f['version']} ({f['kind']}): {f['text']}" for f in fb] or ["- none"]
    mocks = [a["asset_id"] for a in assets if a["mock_approval"]]
    if mocks:
        rep += ["", "## Mock placeholders approved during testing (NOT exported)"] + [f"- {m}" for m in mocks]
    rep += ["", "Human visual review is required for likeness, anatomy, equipment, text and shields; automated QA does not verify these."]
    (out / "review_and_cost_report.md").write_text("\n".join(rep))

    leaks = secrets_leak_scan([p for p in out.rglob("*") if p.suffix in (".json", ".csv", ".txt", ".md")])
    if leaks:
        shutil.rmtree(out)
        raise RuntimeError(f"Export aborted: possible credential found in {leaks}")

    zpath = config.EXPORT_DIR / f"{out.name}.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for p in out.rglob("*"):
            z.write(p, p.relative_to(out))
    state.log("export", f"{zpath.name}: {len(exported)} approved PNG(s)")
    return {"zip": str(zpath.relative_to(config.ROOT)), "folder": str(out.relative_to(config.ROOT)),
            "approved_pngs": exported, "custom_pngs": custom, "mock_approvals_not_exported": [a["asset_id"] for a in assets if a["mock_approval"]],
            "remaining": [a["asset_id"] for a in assets_catalogue if a["status"] == "MISSING" and a["studio_status"] != "APPROVED"]}
