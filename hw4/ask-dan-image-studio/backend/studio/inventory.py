"""Phase 0 inventory: read the six inputs (read-only), classify every catalogue product into a
contextual Dan category independently, detect existing Dan images, prepare references, and
reconcile against the previous audit. No network calls, no paid APIs."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from PIL import Image, ImageChops

from . import config
from .concepts import ASSET_CONCEPTS
from .safezip import safe_extract

# --- Input discovery ----------------------------------------------------------------------


def find_inputs() -> dict[str, Path | None]:
    found: dict[str, Path | None] = {}
    for key, patterns in config.INPUT_PATTERNS.items():
        hit = None
        for pat in patterns:
            matches = sorted(config.INPUT_DIR.glob(pat), key=lambda p: p.stat().st_mtime, reverse=True)
            if matches:
                hit = matches[0]  # newest wins if the user supplies a later version
                break
        found[key] = hit
    return found


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(1 << 20):
            h.update(chunk)
    return h.hexdigest()


# --- Catalogue (read-only) ------------------------------------------------------------------


def load_catalogue(catalogue_zip: Path) -> tuple[list[dict], set[str]]:
    dest = config.WORK_DIR / "catalogue"
    # Only the DB and product photos. Never extract test-account notes or secrets.
    safe_extract(catalogue_zip, dest, only=(".db", ".jpg", ".jpeg", ".png"))
    dbs = list(dest.rglob("campus_customs.db"))
    if not dbs:
        raise FileNotFoundError("campus_customs.db not found in catalogue ZIP")
    db = dbs[0]
    conn = sqlite3.connect(f"{db.resolve().as_uri()}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    try:
        rows = [dict(r) for r in conn.execute(
            "SELECT product_id, name, garment_type, description, search_tags, image_file_path, price FROM catalogue ORDER BY product_id")]
    finally:
        conn.close()
    photos = {p.name for p in db.parent.rglob("*") if p.suffix.lower() in (".jpg", ".jpeg", ".png")}
    for r in rows:
        r["photo_file"] = Path(r["image_file_path"]).name
        r["photo_present"] = r["photo_file"] in photos
        r["photo_rel"] = str((db.parent / r["image_file_path"]).relative_to(config.ROOT))
    return rows, photos


# --- Independent product → Dan category classifier ----------------------------------------

COLLEGES = ["Benjamin Franklin", "Berkeley", "Branford", "Davenport", "Ezra Stiles", "Grace Hopper",
            "Jonathan Edwards", "Morse", "Pauli Murray", "Pierson", "Saybrook", "Silliman",
            "Timothy Dwight", "Trumbull"]
SCHOOLS = {  # slug words → display
    "architecture": "Architecture", "art": "Art", "divinity": "Divinity", "engineering": "Engineering & Applied Science",
    "environment": "Environment", "forest": "Environment", "law": "Law", "management": "Management",
    "medicine": "Medicine", "music": "Music", "nursing": "Nursing", "public-health": "Public Health",
}
FAMILY = ["aunt", "brother", "cousin", "dad", "grandma", "grandpa", "mom", "uncle"]
SPORTS = [  # (slug token(s), display) — order matters: first match wins
    ("field-hockey", "Field Hockey"), ("baseball", "Baseball"), ("basketball", "Basketball"),
    ("football", "Football"), ("yale-bowl", "Football"), ("hockey", "Ice Hockey"), ("lacrosse", "Lacrosse"),
    ("sailing", "Sailing"), ("soccer", "Soccer"), ("squash", "Squash"), ("tennis", "Tennis"),
    ("track", "Track & Field"), ("diving", "Diving"), ("swimming", "Swimming"), ("fencing", "Fencing"),
    ("golf", "Golf"), ("volleyball", "Volleyball"),
]


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def classify(p: dict) -> tuple[str, str, str]:
    """Return (group, category, reason). Pure function of the catalogue row."""
    pid, name, g = p["product_id"], p["name"], p["garment_type"].lower()
    words = f"-{pid}-"
    if "harvard" in pid:
        return "Rivalries", "Harvard", "name mentions Harvard (The Game)"
    for c in COLLEGES:
        if f"-{c.lower().replace(' ', '-')}-" in words:
            return "Colleges", c, f"residential college name '{c}' in product id"
    if "school-of-" in pid or "law-school" in pid or "divinity-school" in pid or "forest-school" in pid:
        for key, disp in SCHOOLS.items():
            if f"-{key}-" in words or f"-{key}" in words:
                return "Schools", disp, f"Yale school '{disp}' in product id"
    for f in FAMILY:
        if f"-{f}-" in words:
            return "Family", f.capitalize(), f"family gift word '{f}' in product id"
    if pid.startswith("crew-left-chest"):
        return "Sports", "Rowing", "'Crew' left-chest design = Yale crew (rowing), not a crewneck"
    for token, disp in SPORTS:
        if f"-{token}-" in words or words.endswith(f"-{token}-"):
            return "Sports", disp, f"sport '{token}' in product id"
    # Plain apparel by garment type (note: 'sailor bulldog' is not sailing; 'raglan crew' is a crewneck)
    if "performance" in g:
        return "Apparel", "Performance Apparel", f"garment_type '{p['garment_type']}'"
    if "jacket" in g and "fleece" not in g:
        return "Apparel", "Jackets / Full-Zips", f"garment_type '{p['garment_type']}'"
    if "hood" in g:
        return "Apparel", "Hoodies", f"garment_type '{p['garment_type']}'"
    if "crew" in g:
        return "Apparel", "Crewnecks", f"garment_type '{p['garment_type']}'"
    if "t-shirt" in g:
        return "Apparel", "T-Shirts", f"garment_type '{p['garment_type']}'"
    return "Apparel", "Other Apparel", f"garment_type '{p['garment_type']}' fits no main apparel group"


def asset_id_for(group: str, category: str) -> str:
    if group == "Apparel":
        return {"Crewnecks": "dan_crewneck", "Hoodies": "dan_hoodie", "Jackets / Full-Zips": "dan_jacket",
                "Other Apparel": "dan_apparel_other", "Performance Apparel": "dan_performance",
                "T-Shirts": "dan_tshirt"}[category]
    prefix = {"Family": "dan_family_", "Schools": "dan_school_", "Sports": "dan_sport_",
              "Colleges": "dan_college_", "Rivalries": "dan_rival_"}[group]
    slug = _slug(category.replace("&", "and"))
    return prefix + slug


# --- Existing Dan images ---------------------------------------------------------------------

EXISTING_RULES = [  # (asset_id, filename regex)
    ("dan_rival_harvard", r"^dan (ivy )?harvard side ?eye"),
    ("dan_sport_baseball", r"^dan baseball"),
    ("dan_sport_basketball", r"^dan basketball"),
    ("dan_sport_football", r"^dan foot ?ball"),
    ("dan_sport_ice_hockey", r"^dan ice hockey"),
    ("dan_sport_rowing", r"^dan (harvard )?crew rowing"),
    ("dan_sport_soccer", r"^dan soccer"),
    ("dan_sport_swimming", r"^dan swimming"),
] + [(f"dan_college_{_slug(c)}", rf"^dan {c.lower().replace('berkeley', 'berke(?:le)?ly')}( college)?\.png$")
     for c in COLLEGES] + [
    (f"dan_rival_{s}", rf"^dan ivy {s}") for s in ("princeton", "upenn", "brown", "columbia", "cornell", "dartmouth")
]


@dataclass
class DanImage:
    filename: str
    rel_path: str
    size: tuple[int, int]
    asset_id: str | None
    disposition: str  # mapped | out_of_scope | collage_excluded | deprecated_dan2 | unmatched
    notes: list[str] = field(default_factory=list)


def scan_existing(dan_zip: Path, needed: set[str]) -> list[DanImage]:
    dest = config.WORK_DIR / "dan_final"
    files = safe_extract(dan_zip, dest, only=(".png", ".jpg", ".jpeg"))
    out = []
    for f in sorted(files):
        name = f.name
        low = name.lower()
        with Image.open(f) as im:
            size = im.size
        rel = str(f.relative_to(config.ROOT))
        if low.startswith("dan2"):
            out.append(DanImage(name, rel, size, None, "deprecated_dan2", ["Deprecated Dan2 look: never a reference"]))
            continue
        if "socks and hoodie" in low or size[0] / size[1] > 2:
            out.append(DanImage(name, rel, size, None, "collage_excluded", ["Two-panel collage: not an individual asset"]))
            continue
        aid = next((a for a, rx in EXISTING_RULES if re.search(rx, low)), None)
        if aid is None:
            out.append(DanImage(name, rel, size, None, "unmatched", ["No catalogue category matches this file"]))
            continue
        notes = []
        if "berkely" in low:
            notes.append("Filename misspells Berkeley (kept as-is; mapped, not renamed)")
        if min(size) < 1024:
            notes.append(f"Below 1024px on one side ({size[0]}×{size[1]})")
        if size[0] != size[1]:
            notes.append("Not square")
        if aid.startswith("dan_college_"):
            notes.append("College shield/lettering should be visually checked against references/shields")
        out.append(DanImage(name, rel, size, aid, "mapped" if aid in needed else "out_of_scope", notes))
    return out


# --- Shields (from the labeled reference screenshots) ----------------------------------------

# Screenshot timestamp → shield identity. Labels are printed under each shield in the source
# images; this table was read from them and is shown in the dashboard for confirmation.
SHIELD_BY_TIME = {
    "9.16.06": ("yale_university", "Yale University"), "9.16.10": ("yale_college", "Yale College"),
    "9.16.15": ("graduate_school_of_arts_and_sciences", "Graduate School of Arts and Sciences"),
    "9.16.20": ("faculty_of_arts_and_sciences", "Faculty of Arts and Sciences"),
    "9.16.25": ("school_of_architecture", "School of Architecture"), "9.16.29": ("school_of_art", "School of Art"),
    "9.16.34": ("divinity_school", "Divinity School"), "9.16.39": ("school_of_drama", "David Geffen School of Drama"),
    "9.16.45": ("school_of_engineering_and_applied_science", "School of Engineering & Applied Science"),
    "9.16.50": ("school_of_the_environment", "School of the Environment"),
    "9.16.55": ("jackson_school_of_global_affairs", "Jackson School of Global Affairs"),
    "9.16.59": ("law_school", "Law School"), "9.17.03": ("school_of_management", "School of Management"),
    "9.17.07": ("school_of_medicine", "School of Medicine"), "9.17.10": ("school_of_music", "School of Music"),
    "9.17.14": ("school_of_nursing", "School of Nursing"), "9.17.18": ("school_of_public_health", "School of Public Health"),
    "9.17.23": ("berkeley_college", "Berkeley College"), "9.17.26": ("branford_college", "Branford College"),
    "9.17.30": ("davenport_college", "Davenport College"), "9.17.34": ("benjamin_franklin_college", "Benjamin Franklin College"),
    "9.17.37": ("grace_hopper_college", "Grace Hopper College"), "9.17.41": ("jonathan_edwards_college", "Jonathan Edwards College"),
    "9.17.45": ("morse_college", "Morse College"), "9.17.49": ("pauli_murray_college", "Pauli Murray College"),
    "9.17.53": ("pierson_college", "Pierson College"), "9.17.57": ("saybrook_college", "Saybrook College"),
    "9.18.00": ("silliman_college", "Silliman College"), "9.18.04": ("ezra_stiles_college", "Ezra Stiles College"),
    "9.18.08": ("timothy_dwight_college", "Timothy Dwight College"), "9.18.11": ("trumbull_college", "Trumbull College"),
}


def _crop_shield(src: Path) -> Image.Image:
    """Keep the shield, drop the printed label underneath; white → transparent."""
    im = Image.open(src).convert("RGB")
    gray = im.convert("L")
    w, h = im.size
    px = gray.load()
    ink_rows = [any(px[x, y] < 235 for x in range(0, w, 2)) for y in range(h)]
    blocks, start = [], None
    for y, ink in enumerate(ink_rows + [False]):
        if ink and start is None:
            start = y
        elif not ink and start is not None:
            blocks.append((start, y))
            start = None
    merged: list[list[int]] = []
    for a, b in blocks:  # merge blocks separated by small gaps (within the shield art)
        if merged and a - merged[-1][1] < max(6, h // 60):
            merged[-1][1] = b
        else:
            merged.append([a, b])
    top, bottom = max(merged, key=lambda ab: ab[1] - ab[0])  # tallest block = the shield
    shield = im.crop((0, top, w, bottom))
    bbox = ImageChops.difference(shield, Image.new("RGB", shield.size, "white")).getbbox()
    if bbox:
        shield = shield.crop(bbox)
    # Background → transparent: ONLY near-white pixels connected to the image border (flood fill).
    # White fields INSIDE the shield (Trumbull, Morse, Davenport…) stay opaque.
    from PIL import ImageDraw
    pad = Image.new("RGB", (shield.width + 2, shield.height + 2), "white")
    pad.paste(shield, (1, 1))
    gray = pad.convert("L")
    mask = gray.point(lambda v: 255 if v > 240 else 0)
    ImageDraw.floodfill(mask, (0, 0), 128)
    bg = mask.point(lambda v: 255 if v == 128 else 0).crop((1, 1, shield.width + 1, shield.height + 1))
    rgba = shield.convert("RGBA")
    rgba.putalpha(ImageChops.invert(bg))
    return rgba


def prepare_references(refs_zip: Path, dan_files: list[DanImage]) -> dict:
    dest = config.WORK_DIR / "references_zip"
    files = safe_extract(refs_zip, dest, only=(".png", ".jpg", ".jpeg"))
    shields, dan_refs, unmapped = [], [], []
    for f in sorted(files):
        m = re.search(r"(\d{1,2}\.\d{2}\.\d{2})", f.name)
        if m and m.group(1) in SHIELD_BY_TIME:
            slug, label = SHIELD_BY_TIME[m.group(1)]
            target = config.REF_SHIELD_DIR / f"{slug}.png"
            if not target.exists():
                _crop_shield(f).save(target)
            shields.append({"slug": slug, "label": label, "path": str(target.relative_to(config.ROOT)),
                            "source": f.name, "mapping": "read from label printed in source image"})
        elif f.name.lower().startswith("welcome dan"):
            dan_refs.append(f)
        else:
            unmapped.append(f.name)

    # Canonical identity references: approved ORIGINAL Dan images only (never Dan2).
    preferred = ["Welcome Dan.png", "Dan Baseball Babe Ruth.png", "Dan Swimming.png", "Dan Ice Hockey.png",
                 "Dan Ivy Harvard Sideeye.png", "Dan Harvard Sideeye.png", "Dan Crew Rowing.png"]
    by_name = {d.filename: config.ROOT / d.rel_path for d in dan_files if d.disposition in ("mapped", "out_of_scope")}
    by_name.update({f.name: f for f in dan_refs})
    chosen = []
    for name in preferred:
        if name in by_name and not name.lower().startswith("dan2") and len(chosen) < 5:
            target = config.REF_DAN_DIR / name.replace(" ", "_")
            if not target.exists():
                Image.open(by_name[name]).convert("RGB").save(target)
            chosen.append({"name": name, "path": str(target.relative_to(config.ROOT))})
    return {"shields": shields, "dan_references": chosen, "unmapped_reference_files": unmapped}


# --- Previous audit (XLSX + TXT) ---------------------------------------------------------------


def load_audit(xlsx: Path | None, txt: Path | None) -> dict:
    audit = {"product_map": {}, "dashboard": {}, "txt_counts": {}}
    if xlsx:
        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        if "102 Product Mappings" in wb.sheetnames:
            rows = list(wb["102 Product Mappings"].iter_rows(values_only=True))
            hdr = rows[0]
            for r in rows[1:]:
                rec = dict(zip(hdr, r))
                if rec.get("Product ID"):
                    audit["product_map"][rec["Product ID"]] = rec
        if "Dashboard" in wb.sheetnames:
            for r in wb["Dashboard"].iter_rows(values_only=True):
                if r and isinstance(r[0], str) and len(r) > 1 and isinstance(r[1], (int, float)):
                    audit["dashboard"][r[0]] = r[1]
    if txt:
        for line in txt.read_text(encoding="utf-8", errors="replace").splitlines():
            m = re.match(r"^([A-Z][A-Z /]+):\s*(\d+)", line.strip())
            if m:
                audit["txt_counts"][m.group(1).strip()] = int(m.group(2))
    return audit


# --- Build manifest + report -------------------------------------------------------------------


def build() -> dict:
    inputs = find_inputs()
    missing_inputs = [k for k, v in inputs.items() if v is None]
    if inputs["catalogue_zip"] is None or inputs["dan_final_zip"] is None:
        raise FileNotFoundError(f"Required inputs missing: {missing_inputs}")

    products, _ = load_catalogue(inputs["catalogue_zip"])
    by_asset: dict[str, dict] = {}
    for p in products:
        group, cat, reason = classify(p)
        aid = asset_id_for(group, cat)
        p.update({"dan_group": group, "dan_category": cat, "dan_asset_id": aid, "mapping_reason": reason})
        a = by_asset.setdefault(aid, {"asset_id": aid, "group": group, "category": cat, "product_ids": []})
        a["product_ids"].append(p["product_id"])

    dan_files = scan_existing(inputs["dan_final_zip"], set(by_asset))
    refs = prepare_references(inputs["dan_refs_zip"], dan_files) if inputs["dan_refs_zip"] else \
        {"shields": [], "dan_references": [], "unmapped_reference_files": []}
    shield_paths = {s["slug"]: s["path"] for s in refs["shields"]}

    existing_by_asset: dict[str, list[str]] = defaultdict(list)
    for d in dan_files:
        if d.asset_id:
            existing_by_asset[d.asset_id].append(d.filename)

    assets = []
    for aid, a in sorted(by_asset.items(), key=lambda kv: (kv[1]["group"], kv[0])):
        concept = ASSET_CONCEPTS.get(aid, {})
        existing = existing_by_asset.get(aid, [])
        shield_slug = concept.get("shield")
        status = "EXISTING_REVIEWED" if existing else ("MISSING" if concept else "OUT_OF_SCOPE")
        qa = []
        if shield_slug and shield_slug not in shield_paths:
            qa.append(f"Shield '{shield_slug}' not supplied: omit heraldry or mark for review")
        assets.append({
            **a,
            "product_count": len(a["product_ids"]),
            "display_name": concept.get("display_name", a["category"]),
            "canonical_filename": f"{aid}.png",
            "existing_files": existing,
            "status": status,
            "pilot": aid in config.PILOT_ASSETS,
            "shield_slug": shield_slug,
            "shield_path": shield_paths.get(shield_slug) if shield_slug else None,
            "setting": concept.get("setting"), "outfit": concept.get("outfit"), "pose": concept.get("pose"),
            "humor": concept.get("humor"), "permitted_text": concept.get("permitted_text", []),
            "qa_warnings": qa,
        })

    # Existing Dan files for categories the catalogue doesn't need → out of scope, kept in inventory.
    for d in dan_files:
        if d.disposition == "out_of_scope" and d.asset_id not in by_asset:
            pass

    audit = load_audit(inputs["audit_xlsx"], inputs["checklist_txt"])
    discrepancies = []
    for p in products:
        prev = audit["product_map"].get(p["product_id"])
        if prev is None:
            discrepancies.append({"product_id": p["product_id"], "issue": "not in previous audit"})
            continue
        if (prev.get("Dan group"), prev.get("Dan category")) != (p["dan_group"], p["dan_category"]):
            discrepancies.append({"product_id": p["product_id"], "issue": "category differs",
                                  "audit": f"{prev.get('Dan group')}: {prev.get('Dan category')}",
                                  "studio": f"{p['dan_group']}: {p['dan_category']}"})
        prev_file = prev.get("Existing Dan file(s)")
        if prev_file:
            for fn in [x.strip() for x in str(prev_file).split(";")]:
                if fn and fn not in existing_by_asset.get(p["dan_asset_id"], []):
                    discrepancies.append({"product_id": p["product_id"], "issue": "audit filename not found in archive",
                                          "audit": fn, "studio": "; ".join(existing_by_asset.get(p["dan_asset_id"], [])) or "-"})
    for pid in set(audit["product_map"]) - {p["product_id"] for p in products}:
        discrepancies.append({"product_id": pid, "issue": "in audit but not in catalogue"})
    # De-duplicate filename issues (one per filename)
    seen, dedup = set(), []
    for d in discrepancies:
        key = (d["issue"], d.get("audit"), d.get("studio")) if d["issue"].startswith("audit filename") else (d["product_id"], d["issue"])
        if key not in seen:
            seen.add(key)
            dedup.append(d)

    in_scope = [a for a in assets if a["status"] in ("MISSING",)]
    counts = {
        "catalogue_products": len(products),
        "product_photos_present": sum(p["photo_present"] for p in products),
        "dan_assets_required": len(assets),
        "existing_assets": sum(a["status"] == "EXISTING_REVIEWED" for a in assets),
        "missing_assets": len(in_scope),
        "dan_files_in_archive": len(dan_files),
        "dan_files_out_of_scope": sum(d.disposition == "out_of_scope" for d in dan_files),
        "collages_excluded": sum(d.disposition == "collage_excluded" for d in dan_files),
        "deprecated_dan2": sum(d.disposition == "deprecated_dan2" for d in dan_files),
        "shields_available": len(refs["shields"]),
        "dan_references": len(refs["dan_references"]),
    }
    previous = {"catalogue_products": audit["dashboard"].get("Catalogue products"),
                "dan_assets_required": audit["txt_counts"].get("DISTINCT DAN CATEGORY ASSETS REQUIRED"),
                "existing_assets": audit["txt_counts"].get("EXISTING DAN CATEGORY ASSETS"),
                "missing_assets": audit["txt_counts"].get("DAN CATEGORY ASSETS TO GENERATE")}

    manifest = {
        "version": 1,
        "inputs": {k: ({"file": v.name, "sha256": sha256(v)} if v else None) for k, v in inputs.items()},
        "missing_inputs": missing_inputs,
        "counts": counts,
        "previous_audit_counts": previous,
        "discrepancies": dedup,
        "assets": assets,
        "products": [{k: p[k] for k in ("product_id", "name", "garment_type", "price", "photo_file", "photo_present",
                                        "photo_rel", "dan_group", "dan_category", "dan_asset_id", "mapping_reason")}
                     for p in products],
        "dan_files": [d.__dict__ for d in dan_files],
        "references": refs,
    }
    tmp = config.MANIFEST_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(manifest, indent=2))
    tmp.replace(config.MANIFEST_PATH)
    write_report(manifest)
    return manifest


def write_report(m: dict) -> Path:
    c, prev = m["counts"], m["previous_audit_counts"]
    lines = ["# Ask Dan Image Studio — Phase 0 inventory (dry run, no paid calls)", "",
             "## Inputs"]
    for k, v in m["inputs"].items():
        lines.append(f"- **{k}**: {v['file'] + ' (sha256 ' + v['sha256'][:12] + '…)' if v else 'MISSING'}")
    lines += ["", "## Reconciled counts", "| Metric | Studio (verified) | Previous audit |", "|---|---|---|"]
    for key in ("catalogue_products", "dan_assets_required", "existing_assets", "missing_assets"):
        lines.append(f"| {key} | {c[key]} | {prev.get(key)} |")
    lines += [f"| product_photos_present | {c['product_photos_present']}/{c['catalogue_products']} | |",
              f"| shields_available | {c['shields_available']} | |", f"| dan_references | {c['dan_references']} | |",
              "", "## Discrepancies vs previous audit"]
    lines += [f"- `{d['product_id']}` — {d['issue']}: audit `{d.get('audit', '')}` → studio `{d.get('studio', '')}`"
              for d in m["discrepancies"]] or ["- None"]
    lines += ["", "## Dan archive files"]
    lines += [f"- {d['filename']} — {d['disposition']}" + (f" → `{d['asset_id']}`" if d["asset_id"] else "")
              + (f" ({'; '.join(d['notes'])})" if d["notes"] else "") for d in m["dan_files"]]
    lines += ["", "## Generation queue (MISSING, catalogue-scoped)"]
    lines += [f"- {'**PILOT** ' if a['pilot'] else ''}`{a['canonical_filename']}` — {a['group']}: {a['category']} "
              f"({a['product_count']} product(s)); shield: {a['shield_slug'] or '—'}"
              for a in m["assets"] if a["status"] == "MISSING"]
    p = config.REPORTS_DIR / "inventory_report.md"
    p.write_text("\n".join(lines) + "\n")
    with open(config.REPORTS_DIR / "product_mapping.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(m["products"][0].keys()))
        w.writeheader()
        w.writerows(m["products"])
    return p
